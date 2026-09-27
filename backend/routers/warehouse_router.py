import io
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from sqlalchemy.orm import Session

import models
import schemas
import serializers
import seed
from auth import get_current_user, log_action, require_roles
from database import get_db

router = APIRouter(prefix="/api/warehouse", tags=["warehouse"])

STOCK_CORRECTION_DISTRIBUTOR_NAME = "Sarrdesign - Weekly Stock Correction (Internal)"
STOCK_CORRECTION_SHOWROOM_NAME = "Weekly Stock Correction"

# Placeholder distributors whose items are "ours to adjust" when reconciling stock -
# the original one-time opening-balance import, plus this recurring weekly one.
# Real distributor/showroom activity is never touched by the stock-correction import.
INTERNAL_STOCK_DISTRIBUTOR_NAMES = {
    seed.OPENING_STOCK_DISTRIBUTOR_NAME,
    STOCK_CORRECTION_DISTRIBUTOR_NAME,
}


def recompute_item_status(item: models.SampleRequestItem):
    """Re-derive an item's high-level status from its current quantity counters.
    Called after anything that changes qty_assigned_showroom (new assignment or a return)."""
    if item.status == models.ItemStatus.REJECTED:
        return
    if item.qty_assigned_showroom <= 0:
        item.status = (models.ItemStatus.FULLY_RECEIVED if item.qty_received_warehouse >= item.qty_approved
                       else models.ItemStatus.PARTIALLY_RECEIVED)
    elif item.qty_assigned_showroom < item.qty_received_warehouse:
        item.status = models.ItemStatus.PARTIALLY_ASSIGNED
    else:
        item.status = models.ItemStatus.ASSIGNED


def recompute_request_status(req: models.SampleRequest):
    item_statuses = {it.status for it in req.items if it.status != models.ItemStatus.REJECTED}
    if not item_statuses:
        return
    if item_statuses.issubset({models.ItemStatus.ASSIGNED}):
        req.status = models.RequestStatus.COMPLETED
    elif models.ItemStatus.ASSIGNED in item_statuses or models.ItemStatus.PARTIALLY_ASSIGNED in item_statuses:
        req.status = models.RequestStatus.PARTIALLY_ASSIGNED
    elif item_statuses.issubset({models.ItemStatus.FULLY_RECEIVED}):
        req.status = models.RequestStatus.FULLY_RECEIVED
    elif models.ItemStatus.PARTIALLY_RECEIVED in item_statuses or models.ItemStatus.FULLY_RECEIVED in item_statuses:
        req.status = models.RequestStatus.PARTIALLY_RECEIVED


# ---------------------------------------------------------------------------
# Receiving (handles: batches never arrive together, delays up to months)
# ---------------------------------------------------------------------------

@router.get("/incoming")
def incoming(db: Session = Depends(get_db),
             user: models.User = Depends(require_roles(models.Role.WAREHOUSE, models.Role.COMMERCIAL_DIRECTOR, models.Role.FACTORY))):
    """Items shipped by the factory but not yet fully received - and items approved
    but never shipped at all (the >3 month delay problem)."""
    statuses = [models.ItemStatus.SENT_TO_FACTORY, models.ItemStatus.IN_PRODUCTION,
                models.ItemStatus.PARTIALLY_SHIPPED, models.ItemStatus.FULLY_SHIPPED,
                models.ItemStatus.PARTIALLY_RECEIVED]
    items = db.query(models.SampleRequestItem).filter(models.SampleRequestItem.status.in_(statuses)).all()
    out = []
    for i in items:
        pending_receive = i.qty_shipped_from_factory - i.qty_received_warehouse
        never_shipped_days = None
        if i.qty_shipped_from_factory == 0:
            never_shipped_days = (models.now() - i.updated_at).days if i.updated_at else None
        out.append({
            **serializers.item_out(i),
            "request_number": i.request.request_number,
            "distributor_name": i.request.distributor.name if i.request.distributor else None,
            "showroom_name": i.request.showroom.name if i.request.showroom else None,
            "pending_receive": pending_receive,
            "outstanding_from_factory": i.qty_approved - i.qty_shipped_from_factory,
            "days_since_last_update": never_shipped_days,
        })
    return out


@router.post("/receipts")
def log_receipt(payload: schemas.WarehouseReceiptIn, db: Session = Depends(get_db),
                 user: models.User = Depends(require_roles(models.Role.WAREHOUSE))):
    item = db.query(models.SampleRequestItem).filter(models.SampleRequestItem.id == payload.item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    if payload.qty_received <= 0:
        raise HTTPException(status_code=400, detail="qty_received must be positive")
    remaining = item.qty_shipped_from_factory - item.qty_received_warehouse
    if payload.qty_received > remaining:
        raise HTTPException(status_code=400, detail=f"Only {remaining} pieces are known to be in transit for this item")

    receipt = models.WarehouseReceipt(
        item_id=item.id, factory_shipment_id=payload.factory_shipment_id,
        qty_received=payload.qty_received, received_by=user.id, notes=payload.notes,
    )
    db.add(receipt)
    item.qty_received_warehouse += payload.qty_received
    item.status = (models.ItemStatus.FULLY_RECEIVED if item.qty_received_warehouse >= item.qty_approved
                   else models.ItemStatus.PARTIALLY_RECEIVED)

    req = item.request
    statuses = {it.status for it in req.items if it.status != models.ItemStatus.REJECTED}
    if statuses and statuses.issubset({models.ItemStatus.FULLY_RECEIVED}):
        req.status = models.RequestStatus.FULLY_RECEIVED
    elif models.ItemStatus.PARTIALLY_RECEIVED in statuses or models.ItemStatus.FULLY_RECEIVED in statuses:
        req.status = models.RequestStatus.PARTIALLY_RECEIVED

    log_action(db, user, "sample_request_item", item.id, "warehouse_receipt", f"received {payload.qty_received}")
    db.commit()
    db.refresh(item)
    return serializers.item_out(item)


@router.post("/shipment-flags")
def raise_shipment_flag(payload: schemas.ShipmentFlagIn, db: Session = Depends(get_db),
                         user: models.User = Depends(require_roles(models.Role.WAREHOUSE))):
    """
    Warehouse raises a flag when what physically arrived doesn't match what the
    factory said it sent (wrong quantity, wrong SKU/color, damage, etc). This never
    blocks receiving or anything downstream - it's purely a visible alert for the
    admin (and factory/director) to look into.
    """
    item = db.query(models.SampleRequestItem).filter(models.SampleRequestItem.id == payload.item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    if not payload.reason or not payload.reason.strip():
        raise HTTPException(status_code=400, detail="Describe the issue")

    flag = models.ShipmentIssueFlag(
        item_id=item.id, factory_shipment_id=payload.factory_shipment_id,
        reason=payload.reason.strip(), raised_by=user.id,
    )
    db.add(flag)
    log_action(db, user, "sample_request_item", item.id, "shipment_flag_raised", payload.reason.strip())
    db.commit()
    db.refresh(flag)
    return serializers.flag_out(flag)


@router.get("/shipment-flags")
def list_shipment_flags(resolved: Optional[bool] = None, db: Session = Depends(get_db),
                         user: models.User = Depends(get_current_user)):
    query = db.query(models.ShipmentIssueFlag)
    if resolved is not None:
        query = query.filter(models.ShipmentIssueFlag.resolved == resolved)
    flags = query.order_by(models.ShipmentIssueFlag.id.desc()).all()
    return [serializers.flag_out(f) for f in flags]


@router.post("/shipment-flags/{flag_id}/resolve")
def resolve_shipment_flag(flag_id: int, payload: schemas.FlagResolveIn, db: Session = Depends(get_db),
                           user: models.User = Depends(require_roles(models.Role.ADMIN))):
    flag = db.query(models.ShipmentIssueFlag).filter(models.ShipmentIssueFlag.id == flag_id).first()
    if not flag:
        raise HTTPException(status_code=404, detail="Flag not found")
    if flag.resolved:
        raise HTTPException(status_code=400, detail="Already resolved")
    flag.resolved = True
    flag.resolved_by = user.id
    flag.resolved_at = models.now()
    flag.resolution_notes = payload.resolution_notes
    log_action(db, user, "sample_request_item", flag.item_id, "shipment_flag_resolved", payload.resolution_notes or "")
    db.commit()
    db.refresh(flag)
    return serializers.flag_out(flag)


@router.get("/stock")
def stock(q: Optional[str] = None, db: Session = Depends(get_db),
          user: models.User = Depends(get_current_user)):
    """Current physical stock sitting in the downtown warehouse, per request item."""
    items = db.query(models.SampleRequestItem).filter(
        models.SampleRequestItem.qty_received_warehouse > 0
    ).all()
    out = []
    for i in items:
        in_wh = i.qty_in_warehouse
        if in_wh <= 0:
            continue
        if q:
            hay = f"{i.product.code} {i.product.family} {i.product.color}".lower()
            if q.lower() not in hay:
                continue
        out.append({
            **serializers.item_out(i),
            "request_number": i.request.request_number,
            "distributor_name": i.request.distributor.name if i.request.distributor else None,
            "showroom_name": i.request.showroom.name if i.request.showroom else None,
        })
    return out


def _next_request_number(db: Session) -> str:
    last = db.query(models.SampleRequest).order_by(models.SampleRequest.id.desc()).first()
    n = 244
    if last and last.request_number:
        try:
            n = int(last.request_number.split("-")[-1])
        except ValueError:
            n = 244
    return f"SD-{n + 1:04d}"


def _get_or_create_correction_showroom(db: Session, admin: models.User):
    dist = db.query(models.Distributor).filter(
        models.Distributor.name == STOCK_CORRECTION_DISTRIBUTOR_NAME
    ).first()
    if dist:
        showroom = db.query(models.Showroom).filter(
            models.Showroom.distributor_id == dist.id,
            models.Showroom.name == STOCK_CORRECTION_SHOWROOM_NAME,
        ).first()
        if showroom:
            return dist, showroom
    else:
        dist = models.Distributor(
            name=STOCK_CORRECTION_DISTRIBUTOR_NAME,
            type=models.PartnerType.DISTRIBUTOR,
            brand="Sarrdesign",
            notes=(
                "Internal placeholder, not a real distributor. Holds the running "
                "per-SKU correction used to reconcile Warehouse Stock against the "
                "team's weekly physical stock-count file. Updated in place each "
                "week - it does not accumulate a separate row per week."
            ),
        )
        db.add(dist)
        db.flush()
    showroom = models.Showroom(
        distributor_id=dist.id,
        name=STOCK_CORRECTION_SHOWROOM_NAME,
        address="N/A - placeholder location for weekly stock reconciliation, not a physical showroom",
        min_order_qty=0,
    )
    db.add(showroom)
    db.flush()
    return dist, showroom


@router.post("/stock-correction/import")
def import_stock_correction(file: UploadFile = File(...), db: Session = Depends(get_db),
                             user: models.User = Depends(require_roles(models.Role.ADMIN))):
    """
    Reconcile Warehouse Stock against the team's weekly stock-count Excel file
    (same format as the original opening-stock ledger: a 'Samples request'-like
    sheet with code in column C and quantity in column F, header rows first).

    For every recognized SKU, the app's current total warehouse stock for that
    SKU (summed across every request item for that product) is corrected to
    match the quantity in the file - a blank/empty quantity cell means zero.
    The correction is applied by adjusting a single dedicated per-SKU entry
    (under an internal 'Weekly Stock Correction' placeholder) so re-running
    this import replaces last week's correction rather than stacking a new one.
    """
    content = file.file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read this file as an Excel (.xlsx) workbook")

    ws = None
    for name in wb.sheetnames:
        if name.strip().lower() == "samples request":
            ws = wb[name]
            break
    if ws is None:
        for name in wb.sheetnames:
            if "sample" in name.strip().lower():
                ws = wb[name]
                break
    if ws is None:
        raise HTTPException(
            status_code=400,
            detail=f"No 'Samples request' sheet found. Sheets in this file: {', '.join(wb.sheetnames)}",
        )

    # code = column C (3), qty = column F (6); rows 1-2 are headers/section labels.
    # Real SKU codes all start with "SD" - filters out family/section header rows
    # (e.g. a family name in Arabic sitting alone in column C).
    file_data = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or len(row) < 6:
            continue
        raw_code = row[2]
        if raw_code is None:
            continue
        code = str(raw_code).strip()
        if not code.upper().startswith("SD"):
            continue
        qty = row[5]
        qty = int(qty) if isinstance(qty, (int, float)) else 0
        file_data[code] = qty

    if not file_data:
        raise HTTPException(status_code=400, detail="No SKU rows found in this file (expected codes starting with 'SD' in column C)")

    admin_user = db.query(models.User).filter(models.User.role == models.Role.ADMIN).first() or user
    _, correction_showroom = _get_or_create_correction_showroom(db, admin_user)

    updated, created, unchanged = [], [], []
    unmatched = []
    capped = []

    for code, target in sorted(file_data.items()):
        product = db.query(models.Product).filter(models.Product.code == code).first()
        if not product:
            if target > 0:
                unmatched.append({"code": code, "qty": target})
            continue

        all_items = db.query(models.SampleRequestItem).filter(
            models.SampleRequestItem.product_id == product.id
        ).all()
        current_total = sum(i.qty_in_warehouse for i in all_items)

        # "Ours to adjust": any item sitting under the opening-balance or weekly-
        # correction placeholders. Everything else (real distributor/showroom
        # activity) is left alone. If more than one such item exists for this SKU
        # (e.g. both an opening-balance entry and an older correction entry),
        # they're consolidated into the first one found and the rest zeroed out,
        # so a SKU only ever carries one internal placeholder line going forward.
        internal_items = [
            i for i in all_items
            if i.request.distributor and i.request.distributor.name in INTERNAL_STOCK_DISTRIBUTOR_NAMES
        ]
        internal_total = sum(i.qty_in_warehouse for i in internal_items)
        other_total = current_total - internal_total

        desired_internal_val = target - other_total
        new_internal_val = max(0, desired_internal_val)
        if desired_internal_val < 0:
            # Real (non-internal) stock alone already exceeds this week's count -
            # can't reduce below what isn't ours to touch.
            capped.append({
                "code": code, "target": target, "reachable_total": other_total + new_internal_val,
            })

        primary = internal_items[0] if internal_items else None
        extra_internal = internal_items[1:]

        if primary:
            old_val = internal_total
            if new_internal_val != old_val:
                delta = new_internal_val - primary.qty_in_warehouse
                if delta > 0:
                    db.add(models.WarehouseReceipt(
                        item_id=primary.id, qty_received=delta, received_by=admin_user.id,
                        notes=f"Weekly stock correction ({file.filename})",
                    ))
                primary.qty_received_warehouse = primary.qty_received_warehouse + delta
                primary.qty_requested = primary.qty_received_warehouse
                primary.qty_approved = primary.qty_received_warehouse
                primary.status = models.ItemStatus.FULLY_RECEIVED
                for extra in extra_internal:
                    if extra.qty_in_warehouse:
                        extra.qty_received_warehouse = extra.qty_assigned_showroom
                        extra.status = models.ItemStatus.FULLY_RECEIVED
                log_action(db, user, "product", product.id, "stock_correction",
                           f"code={code} {old_val} -> {new_internal_val} (file target {target})")
                updated.append({"code": code, "previous_total": current_total, "new_total": other_total + new_internal_val, "delta": new_internal_val - old_val})
            else:
                unchanged.append({"code": code, "total": current_total})
        elif new_internal_val > 0:
            req = models.SampleRequest(
                request_number=_next_request_number(db),
                sales_manager_id=admin_user.id,
                distributor_id=correction_showroom.distributor_id,
                showroom_id=correction_showroom.id,
                purpose="stock_correction",
                distributor_order_ref=f"Weekly stock correction import ({file.filename})",
                status=models.RequestStatus.FULLY_RECEIVED,
                notes="Auto-managed weekly stock correction entry - not a real distributor request.",
            )
            db.add(req)
            db.flush()
            item = models.SampleRequestItem(
                request_id=req.id, product_id=product.id, qty_requested=new_internal_val,
                is_dummy=product.default_dummy if product.default_dummy is not None else True,
                sku_matches_stock=True, min_order_met=True, qty_approved=new_internal_val,
                status=models.ItemStatus.FULLY_RECEIVED,
                qty_shipped_from_factory=new_internal_val, qty_received_warehouse=new_internal_val,
            )
            db.add(item)
            db.flush()
            shipment = models.FactoryShipment(
                item_id=item.id, qty_sent=new_internal_val, batch_ref="STOCK-CORRECTION", created_by=admin_user.id,
            )
            db.add(shipment)
            db.flush()
            db.add(models.WarehouseReceipt(
                item_id=item.id, factory_shipment_id=shipment.id, qty_received=new_internal_val,
                received_by=admin_user.id, notes=f"Weekly stock correction ({file.filename})",
            ))
            log_action(db, user, "product", product.id, "stock_correction",
                       f"code={code} 0 -> {new_internal_val} (file target {target}, new entry)")
            created.append({"code": code, "previous_total": current_total, "new_total": other_total + new_internal_val})
        else:
            unchanged.append({"code": code, "total": current_total})

    log_action(db, user, "product", 0, "stock_correction_import",
               f"file={file.filename} skus_in_file={len(file_data)} updated={len(updated)} "
               f"created={len(created)} unmatched={len(unmatched)} capped={len(capped)}")
    db.commit()

    return {
        "skus_in_file": len(file_data),
        "updated": updated,
        "created": created,
        "unchanged_count": len(unchanged),
        "unmatched": unmatched,
        "capped": capped,
    }


# ---------------------------------------------------------------------------
# Dismissal approvals - warehouse/assembly can SEE stock but need admin sign-off
# to move anything out (new assignment OR reassignment between traders)
# ---------------------------------------------------------------------------

@router.post("/dismissal-requests")
def request_dismissal(payload: schemas.DismissalRequestIn, db: Session = Depends(get_db),
                       user: models.User = Depends(require_roles(models.Role.WAREHOUSE, models.Role.ASSEMBLY))):
    item = db.query(models.SampleRequestItem).filter(models.SampleRequestItem.id == payload.item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    showroom = db.query(models.Showroom).filter(models.Showroom.id == payload.showroom_id).first()
    if not showroom:
        raise HTTPException(status_code=404, detail="Showroom not found")
    if payload.qty > item.qty_in_warehouse and not payload.reassign_from_assignment_id:
        raise HTTPException(status_code=400, detail=f"Only {item.qty_in_warehouse} in warehouse stock for this item")

    dr = models.DismissalApproval(
        item_id=item.id, showroom_id=payload.showroom_id, qty=payload.qty,
        requested_by=user.id, reason=payload.reason,
        reassign_from_assignment_id=payload.reassign_from_assignment_id,
    )
    db.add(dr)
    log_action(db, user, "sample_request_item", item.id, "dismissal_requested",
               f"qty={payload.qty} showroom={payload.showroom_id}")
    db.commit()
    db.refresh(dr)
    return serializers.dismissal_out(dr)


@router.get("/dismissal-requests")
def list_dismissal_requests(status: Optional[str] = "pending", db: Session = Depends(get_db),
                             user: models.User = Depends(get_current_user)):
    query = db.query(models.DismissalApproval)
    if status:
        try:
            query = query.filter(models.DismissalApproval.status == models.DismissalStatus(status))
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid status")
    drs = query.order_by(models.DismissalApproval.id.desc()).all()
    return [serializers.dismissal_out(d) for d in drs]


@router.post("/dismissal-requests/{dismissal_id}/decide")
def decide_dismissal(dismissal_id: int, payload: schemas.DismissalDecisionIn, db: Session = Depends(get_db),
                      user: models.User = Depends(require_roles(models.Role.ADMIN))):
    dr = db.query(models.DismissalApproval).filter(models.DismissalApproval.id == dismissal_id).first()
    if not dr:
        raise HTTPException(status_code=404, detail="Dismissal request not found")
    if dr.status != models.DismissalStatus.PENDING:
        raise HTTPException(status_code=400, detail="Already decided")
    if payload.decision not in ("approve", "reject"):
        raise HTTPException(status_code=400, detail="decision must be approve or reject")
    dr.status = models.DismissalStatus.APPROVED if payload.decision == "approve" else models.DismissalStatus.REJECTED
    dr.decided_by = user.id
    dr.decision_notes = payload.decision_notes
    dr.date_decided = models.now()
    log_action(db, user, "dismissal_approval", dr.id, f"dismissal_{payload.decision}", payload.decision_notes or "")
    db.commit()
    db.refresh(dr)
    return serializers.dismissal_out(dr)


# ---------------------------------------------------------------------------
# Showroom assignment (the assembly team physically displaying the sample)
# ---------------------------------------------------------------------------

@router.get("/showroom/{showroom_id}/current")
def showroom_current(showroom_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """What's currently displayed at a showroom - used to catch 'we gave them 2 of the same sample'."""
    assignments = db.query(models.ShowroomAssignment).filter(
        models.ShowroomAssignment.showroom_id == showroom_id,
        models.ShowroomAssignment.status == models.AssignmentStatus.DISPLAYED,
    ).all()
    return [serializers.assignment_out(a) for a in assignments]


@router.post("/assignments")
def create_assignment(dismissal_id: int, qty: Optional[int] = None, confirm_duplicate: bool = False,
                       db: Session = Depends(get_db),
                       user: models.User = Depends(require_roles(models.Role.ASSEMBLY, models.Role.WAREHOUSE))):
    dr = db.query(models.DismissalApproval).filter(models.DismissalApproval.id == dismissal_id).first()
    if not dr:
        raise HTTPException(status_code=404, detail="Dismissal request not found")
    if dr.status != models.DismissalStatus.APPROVED:
        raise HTTPException(status_code=400, detail="This dismissal has not been approved by the admin yet")

    already_used = db.query(models.ShowroomAssignment).filter(
        models.ShowroomAssignment.dismissal_approval_id == dr.id
    ).first()
    if already_used:
        raise HTTPException(status_code=400, detail="This approval has already been used to create an assignment")

    place_qty = qty or dr.qty
    if place_qty > dr.qty:
        raise HTTPException(status_code=400, detail=f"Approved for at most {dr.qty} pieces")

    # duplicate-display guard: same product code+color already live at this showroom
    duplicate = db.query(models.ShowroomAssignment).join(
        models.SampleRequestItem, models.ShowroomAssignment.item_id == models.SampleRequestItem.id
    ).filter(
        models.ShowroomAssignment.showroom_id == dr.showroom_id,
        models.ShowroomAssignment.status == models.AssignmentStatus.DISPLAYED,
        models.SampleRequestItem.product_id == dr.item.product_id,
    ).first()
    if duplicate and not confirm_duplicate:
        raise HTTPException(
            status_code=409,
            detail=f"This exact sample ({dr.item.product.code} / {dr.item.product.color}) is already displayed "
                   f"at this showroom (assignment #{duplicate.id}, {duplicate.date_assigned.strftime('%Y-%m-%d')}). "
                   f"Resend with confirm_duplicate=true to place it anyway.",
        )

    assignment = models.ShowroomAssignment(
        item_id=dr.item_id, showroom_id=dr.showroom_id, dismissal_approval_id=dr.id,
        qty=place_qty, assigned_by=user.id,
        reassigned_from_id=dr.reassign_from_assignment_id,
    )
    db.add(assignment)

    item = dr.item
    item.qty_assigned_showroom += place_qty
    recompute_item_status(item)

    if dr.reassign_from_assignment_id:
        old = db.query(models.ShowroomAssignment).filter(
            models.ShowroomAssignment.id == dr.reassign_from_assignment_id
        ).first()
        if old:
            old.status = models.AssignmentStatus.REASSIGNED

    recompute_request_status(item.request)

    log_action(db, user, "showroom_assignment", 0, "assign",
               f"item={item.id} showroom={dr.showroom_id} qty={place_qty}")
    db.commit()
    db.refresh(assignment)
    return serializers.assignment_out(assignment)


# ---------------------------------------------------------------------------
# Returns
# ---------------------------------------------------------------------------

@router.post("/returns")
def record_return(payload: schemas.ReturnEventIn, db: Session = Depends(get_db),
                   user: models.User = Depends(require_roles(models.Role.WAREHOUSE, models.Role.ASSEMBLY))):
    assignment = db.query(models.ShowroomAssignment).filter(
        models.ShowroomAssignment.id == payload.assignment_id
    ).first()
    if not assignment:
        raise HTTPException(status_code=404, detail="Assignment not found")
    already_returned = sum(r.qty_returned for r in assignment.returns)
    remaining = assignment.qty - already_returned
    if payload.qty_returned > remaining:
        raise HTTPException(status_code=400, detail=f"Only {remaining} pieces remain displayed for this assignment")
    try:
        reason = models.ReturnReason(payload.reason)
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid return reason")

    ret = models.ReturnEvent(
        assignment_id=assignment.id, qty_returned=payload.qty_returned, reason=reason,
        notes=payload.notes, recorded_by=user.id, restocked=payload.restocked,
    )
    db.add(ret)
    if payload.qty_returned >= remaining:
        assignment.status = models.AssignmentStatus.RETURNED

    item = assignment.item
    item.qty_assigned_showroom -= payload.qty_returned
    if payload.restocked:
        item.qty_returned_warehouse += payload.qty_returned
    else:
        # scrapped / sent for repair - no longer counts as received stock either
        item.qty_received_warehouse -= payload.qty_returned
    recompute_item_status(item)
    recompute_request_status(item.request)

    log_action(db, user, "showroom_assignment", assignment.id, "return",
               f"qty={payload.qty_returned} reason={reason.value}")
    db.commit()
    db.refresh(ret)
    return serializers.return_out(ret)


# ---------------------------------------------------------------------------
# Historical archive - per distributor / per showroom
# ---------------------------------------------------------------------------

@router.get("/history/distributor/{distributor_id}")
def distributor_history(distributor_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    dist = db.query(models.Distributor).filter(models.Distributor.id == distributor_id).first()
    if not dist:
        raise HTTPException(status_code=404, detail="Distributor not found")
    requests = db.query(models.SampleRequest).filter(models.SampleRequest.distributor_id == distributor_id).order_by(models.SampleRequest.id.desc()).all()

    timeline = []
    for r in requests:
        timeline.append({"date": r.date_created.isoformat() if r.date_created else None, "type": "request_submitted",
                          "detail": f"{r.request_number} ({r.showroom.name if r.showroom else ''}) - {len(r.items)} SKUs"})
        for it in r.items:
            for sh in it.shipments:
                timeline.append({"date": sh.date_sent.isoformat() if sh.date_sent else None, "type": "factory_shipment",
                                  "detail": f"{r.request_number} {it.product.code}: shipped {sh.qty_sent}"})
            for rc in it.receipts:
                timeline.append({"date": rc.date_received.isoformat() if rc.date_received else None, "type": "warehouse_receipt",
                                  "detail": f"{r.request_number} {it.product.code}: received {rc.qty_received}"})
            for asg in it.assignments:
                timeline.append({"date": asg.date_assigned.isoformat() if asg.date_assigned else None, "type": "showroom_assignment",
                                  "detail": f"{it.product.code} / {it.product.color}: placed at {asg.showroom.name if asg.showroom else ''} (qty {asg.qty}) [{asg.status.value}]"})
                for ret in asg.returns:
                    timeline.append({"date": ret.date_returned.isoformat() if ret.date_returned else None, "type": "return",
                                      "detail": f"{it.product.code}: returned {ret.qty_returned} from {asg.showroom.name if asg.showroom else ''} - {ret.reason.value}"})

    timeline.sort(key=lambda x: x["date"] or "", reverse=True)
    return {
        "distributor": serializers.distributor_out(dist),
        "showrooms": [serializers.showroom_out(s) for s in dist.showrooms],
        "requests": [serializers.request_out(r, include_items=False) for r in requests],
        "timeline": timeline,
    }


@router.get("/history/showroom/{showroom_id}")
def showroom_history(showroom_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    sr = db.query(models.Showroom).filter(models.Showroom.id == showroom_id).first()
    if not sr:
        raise HTTPException(status_code=404, detail="Showroom not found")
    assignments = db.query(models.ShowroomAssignment).filter(models.ShowroomAssignment.showroom_id == showroom_id).order_by(models.ShowroomAssignment.id.desc()).all()
    currently_displayed = [a for a in assignments if a.status == models.AssignmentStatus.DISPLAYED]
    return {
        "showroom": serializers.showroom_out(sr),
        "currently_displayed": [serializers.assignment_out(a) for a in currently_displayed],
        "all_assignments": [serializers.assignment_out(a) for a in assignments],
        "returns": [serializers.return_out(r) for a in assignments for r in a.returns],
    }
