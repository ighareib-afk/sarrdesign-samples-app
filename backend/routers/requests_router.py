import io
from datetime import datetime
from typing import Optional

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

import models
import schemas
import serializers
from auth import get_current_user, log_action, require_roles
from database import get_db
import os

router = APIRouter(prefix="/api/requests", tags=["requests"])

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)


def next_request_number(db: Session) -> str:
    last = db.query(models.SampleRequest).order_by(models.SampleRequest.id.desc()).first()
    n = 244
    if last and last.request_number:
        try:
            n = int(last.request_number.split("-")[-1])
        except ValueError:
            n = 244
    return f"SD-{n + 1:04d}"


@router.post("")
def create_request(payload: schemas.RequestCreateIn, db: Session = Depends(get_db),
                    user: models.User = Depends(require_roles(models.Role.SALES_MANAGER))):
    if not payload.items:
        raise HTTPException(status_code=400, detail="Add at least one product line")
    showroom = db.query(models.Showroom).filter(models.Showroom.id == payload.showroom_id).first()
    if not showroom or showroom.distributor_id != payload.distributor_id:
        raise HTTPException(status_code=400, detail="Showroom does not belong to that distributor")

    req = models.SampleRequest(
        request_number=next_request_number(db),
        sales_manager_id=user.id,
        distributor_id=payload.distributor_id,
        showroom_id=payload.showroom_id,
        purpose=payload.purpose,
        distributor_order_ref=payload.distributor_order_ref,
        notes=payload.notes,
        status=models.RequestStatus.SUBMITTED,
    )
    db.add(req)
    db.flush()

    min_qty = showroom.min_order_qty or 5
    for line in payload.items:
        product = db.query(models.Product).filter(models.Product.id == line.product_id).first()
        if not product:
            raise HTTPException(status_code=400, detail=f"Product {line.product_id} not found")
        min_met = None
        if line.distributor_order_qty is not None:
            min_met = line.distributor_order_qty >= min_qty
        item = models.SampleRequestItem(
            request_id=req.id, product_id=product.id, qty_requested=line.qty_requested,
            is_dummy=line.is_dummy, distributor_order_qty=line.distributor_order_qty,
            min_order_met=min_met, status=models.ItemStatus.PENDING_REVIEW,
        )
        db.add(item)

    log_action(db, user, "sample_request", req.id, "submit", req.request_number)
    db.add(models.ApprovalAction(request_id=req.id, actor_id=user.id, action="submit", notes="Request submitted"))
    db.commit()
    db.refresh(req)
    return serializers.request_out(req)


@router.get("")
def list_requests(status: Optional[str] = None, mine: bool = False, distributor_id: Optional[int] = None,
                   db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    query = db.query(models.SampleRequest)
    if mine and user.role == models.Role.SALES_MANAGER:
        query = query.filter(models.SampleRequest.sales_manager_id == user.id)
    if status:
        try:
            query = query.filter(models.SampleRequest.status == models.RequestStatus(status))
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid status")
    if distributor_id:
        query = query.filter(models.SampleRequest.distributor_id == distributor_id)
    reqs = query.order_by(models.SampleRequest.id.desc()).all()
    return [serializers.request_out(r, include_items=False) for r in reqs]


@router.get("/{request_id}")
def get_request(request_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    r = db.query(models.SampleRequest).filter(models.SampleRequest.id == request_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Request not found")
    out = serializers.request_out(r)
    out["approvals"] = [serializers.approval_out(a) for a in r.approvals]
    out["attachments"] = [
        {"id": a.id, "file_name": a.file_name, "kind": a.kind,
         "created_at": a.created_at.isoformat() if a.created_at else None}
        for a in r.attachments
    ]
    return out


@router.post("/{request_id}/review")
def review_request(request_id: int, payload: schemas.RequestDecisionIn, db: Session = Depends(get_db),
                    user: models.User = Depends(require_roles(models.Role.COMMERCIAL_DIRECTOR))):
    r = db.query(models.SampleRequest).filter(models.SampleRequest.id == request_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Request not found")
    if r.status not in (models.RequestStatus.SUBMITTED, models.RequestStatus.DIRECTOR_APPROVED):
        raise HTTPException(status_code=400, detail=f"Request is in status {r.status.value}, cannot review")

    items_by_id = {i.id: i for i in r.items}
    for decision in payload.items:
        item = items_by_id.get(decision.item_id)
        if not item:
            raise HTTPException(status_code=400, detail=f"Item {decision.item_id} not on this request")
        item.sku_matches_stock = decision.sku_matches_stock
        item.min_order_met = decision.min_order_met if decision.min_order_met is not None else item.min_order_met
        item.director_notes = decision.director_notes
        if decision.decision == "approve":
            item.status = models.ItemStatus.APPROVED
            item.qty_approved = decision.qty_approved if decision.qty_approved is not None else item.qty_requested
        elif decision.decision == "reject":
            item.status = models.ItemStatus.REJECTED
            item.qty_approved = 0
        else:
            raise HTTPException(status_code=400, detail="decision must be approve or reject")
        db.add(models.ApprovalAction(
            request_id=r.id, item_id=item.id, actor_id=user.id, action=decision.decision,
            notes=decision.director_notes,
        ))

    statuses = {i.status for i in r.items}
    if statuses == {models.ItemStatus.REJECTED}:
        r.status = models.RequestStatus.DIRECTOR_REJECTED
    elif models.ItemStatus.PENDING_REVIEW in statuses:
        r.status = models.RequestStatus.SUBMITTED
    else:
        r.status = models.RequestStatus.DIRECTOR_APPROVED

    log_action(db, user, "sample_request", r.id, "review", payload.notes or "")
    db.commit()
    db.refresh(r)
    return serializers.request_out(r)


@router.post("/items/{item_id}/re-request")
def re_request_item(item_id: int, payload: schemas.ItemReRequestIn, db: Session = Depends(get_db),
                     user: models.User = Depends(get_current_user)):
    """
    Bring a director-rejected line back for review without redoing the whole data
    entry: the sales manager who owns the request (or the admin) can tweak the
    quantity / dummy flag / order qty and resend it to Mostafa - no need to delete
    and recreate the request.
    """
    item = db.query(models.SampleRequestItem).filter(models.SampleRequestItem.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    r = item.request

    is_owner = user.role == models.Role.SALES_MANAGER and r.sales_manager_id == user.id
    if user.role != models.Role.ADMIN and not is_owner:
        raise HTTPException(status_code=403, detail="Only the request's sales manager or an admin can re-request this line")
    if item.status != models.ItemStatus.REJECTED:
        raise HTTPException(status_code=400, detail=f"Item is in status {item.status.value} - only a rejected line can be re-requested")

    if payload.qty_requested is not None:
        if payload.qty_requested <= 0:
            raise HTTPException(status_code=400, detail="qty_requested must be positive")
        item.qty_requested = payload.qty_requested
    if payload.is_dummy is not None:
        item.is_dummy = payload.is_dummy
    if payload.distributor_order_qty is not None:
        item.distributor_order_qty = payload.distributor_order_qty
        min_qty = (r.showroom.min_order_qty or 5) if r.showroom else 5
        item.min_order_met = payload.distributor_order_qty >= min_qty

    # back for a fresh look - clear the previous decision so it doesn't read as
    # already-reviewed, and the director's checks (SKU match / notes) start clean
    item.status = models.ItemStatus.PENDING_REVIEW
    item.qty_approved = 0
    item.sku_matches_stock = None
    item.director_notes = None

    # only recompute the request's own status if it hasn't progressed past review yet
    # (once other lines have gone to the factory, this stays wherever they put it)
    if r.status in (models.RequestStatus.SUBMITTED, models.RequestStatus.DIRECTOR_APPROVED,
                    models.RequestStatus.DIRECTOR_REJECTED):
        statuses = {i.status for i in r.items}
        if statuses == {models.ItemStatus.REJECTED}:
            r.status = models.RequestStatus.DIRECTOR_REJECTED
        elif models.ItemStatus.PENDING_REVIEW in statuses:
            r.status = models.RequestStatus.SUBMITTED
        else:
            r.status = models.RequestStatus.DIRECTOR_APPROVED

    db.add(models.ApprovalAction(
        request_id=r.id, item_id=item.id, actor_id=user.id, action="re_request", notes=payload.notes,
    ))
    log_action(db, user, "sample_request_item", item.id, "re_request", payload.notes or "")
    db.commit()
    db.refresh(item)
    return serializers.item_out(item)


def _item_is_deletable(item: models.SampleRequestItem) -> bool:
    """Safe to delete only before anything has actually left the factory for it -
    past that point, deleting it would leave shipment/stock numbers inconsistent."""
    return item.qty_shipped_from_factory == 0 and len(item.shipments) == 0


@router.delete("/{request_id}/items/{item_id}")
def delete_item(request_id: int, item_id: int, reason: Optional[str] = None, db: Session = Depends(get_db),
                 user: models.User = Depends(require_roles(models.Role.ADMIN))):
    """Admin-only: remove a single wrongly-entered line. Blocked once the factory
    has shipped anything for it. If it's the request's last line, the whole
    (now-empty) request is removed too."""
    r = db.query(models.SampleRequest).filter(models.SampleRequest.id == request_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Request not found")
    item = next((i for i in r.items if i.id == item_id), None)
    if not item:
        raise HTTPException(status_code=404, detail="Item not on this request")
    if not _item_is_deletable(item):
        raise HTTPException(status_code=400, detail="This line has already been shipped from the factory and can no longer be deleted")

    code = item.product.code if item.product else str(item.product_id)
    detail = f"{r.request_number}: deleted line {code} (qty {item.qty_requested})"
    if reason:
        detail += f" - reason: {reason}"
    log_action(db, user, "sample_request_item", item.id, "delete", detail)

    db.query(models.ApprovalAction).filter(models.ApprovalAction.item_id == item.id).delete()
    db.delete(item)
    db.flush()

    remaining = [i for i in r.items if i.id != item_id]
    if not remaining:
        log_action(db, user, "sample_request", r.id, "delete", f"{r.request_number}: deleted (last line removed)")
        db.delete(r)
        db.commit()
        return {"deleted": True, "request_deleted": True}

    if r.status in (models.RequestStatus.SUBMITTED, models.RequestStatus.DIRECTOR_APPROVED,
                    models.RequestStatus.DIRECTOR_REJECTED):
        statuses = {i.status for i in remaining}
        if statuses == {models.ItemStatus.REJECTED}:
            r.status = models.RequestStatus.DIRECTOR_REJECTED
        elif models.ItemStatus.PENDING_REVIEW in statuses:
            r.status = models.RequestStatus.SUBMITTED
        else:
            r.status = models.RequestStatus.DIRECTOR_APPROVED
    db.commit()
    return {"deleted": True, "request_deleted": False}


@router.delete("/{request_id}")
def delete_request(request_id: int, reason: Optional[str] = None, db: Session = Depends(get_db),
                    user: models.User = Depends(require_roles(models.Role.ADMIN))):
    """Admin-only: remove an entire wrongly-entered request. Blocked if any of its
    lines have already been shipped from the factory - delete just the unshipped
    lines individually in that case."""
    r = db.query(models.SampleRequest).filter(models.SampleRequest.id == request_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Request not found")
    undeletable = [i for i in r.items if not _item_is_deletable(i)]
    if undeletable:
        codes = ", ".join(i.product.code if i.product else str(i.id) for i in undeletable)
        raise HTTPException(
            status_code=400,
            detail=f"Cannot delete: already shipped from the factory for: {codes}. Delete the remaining unshipped lines individually instead.",
        )

    detail = f"{r.request_number} deleted ({len(r.items)} line(s))"
    if reason:
        detail += f" - reason: {reason}"
    log_action(db, user, "sample_request", r.id, "delete", detail)
    db.delete(r)
    db.commit()
    return {"deleted": True}


@router.post("/{request_id}/send-to-factory")
def send_to_factory(request_id: int, db: Session = Depends(get_db),
                     user: models.User = Depends(require_roles(models.Role.COMMERCIAL_DIRECTOR))):
    r = db.query(models.SampleRequest).filter(models.SampleRequest.id == request_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Request not found")
    if r.status != models.RequestStatus.DIRECTOR_APPROVED:
        raise HTTPException(status_code=400, detail="Request must be director-approved first")
    for item in r.items:
        if item.status == models.ItemStatus.APPROVED:
            item.status = models.ItemStatus.SENT_TO_FACTORY
    r.status = models.RequestStatus.SENT_TO_FACTORY
    log_action(db, user, "sample_request", r.id, "send_to_factory")
    db.add(models.ApprovalAction(request_id=r.id, actor_id=user.id, action="send_to_factory"))
    db.commit()
    db.refresh(r)
    return serializers.request_out(r)


@router.get("/{request_id}/export.xlsx")
def export_request(request_id: int, db: Session = Depends(get_db),
                    user: models.User = Depends(require_roles(models.Role.FACTORY, models.Role.COMMERCIAL_DIRECTOR))):
    r = db.query(models.SampleRequest).filter(models.SampleRequest.id == request_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Request not found")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sample Production List"

    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws["A1"] = "Sarrdesign - Sample Production & Delivery List"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:H1")
    ws["A2"] = f"Request No: {r.request_number}   |   Date: {r.date_created.strftime('%Y-%m-%d') if r.date_created else ''}"
    ws.merge_cells("A2:H2")
    ws["A3"] = f"Distributor / Trader: {r.distributor.name if r.distributor else ''}"
    ws.merge_cells("A3:H3")
    ws["A4"] = f"Showroom: {r.showroom.name if r.showroom else ''}   |   Location: {r.showroom.address or ''} ({r.showroom.governorate or ''} - {r.showroom.area or ''})"
    ws.merge_cells("A4:H4")

    headers = ["#", "Product Code", "Color", "Family", "Description", "Qty Approved", "Dummy / Sellable", "Notes"]
    ws.append([])
    ws.append(headers)
    header_row = ws.max_row
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=header_row, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")

    idx = 1
    for item in r.items:
        if item.status.value not in ("sent_to_factory", "in_production", "partially_shipped", "fully_shipped",
                                       "partially_received", "fully_received", "assigned", "partially_assigned"):
            continue
        ws.append([
            idx, item.product.code if item.product else "", item.product.color if item.product else "",
            item.product.family if item.product else "",
            item.product.description_en if item.product else "",
            item.qty_approved, "Dummy" if item.is_dummy else "Sellable", item.director_notes or "",
        ])
        idx += 1

    widths = [4, 20, 14, 14, 45, 12, 16, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"{r.request_number}_sample_list.xlsx"
    log_action(db, user, "sample_request", r.id, "export_excel")
    db.commit()
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/{request_id}/attachments")
def upload_attachment(request_id: int, kind: str = Form("distributor_order"), file: UploadFile = File(...),
                       db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    r = db.query(models.SampleRequest).filter(models.SampleRequest.id == request_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Request not found")
    safe_name = f"{request_id}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{file.filename}"
    path = os.path.join(UPLOAD_DIR, safe_name)
    with open(path, "wb") as f:
        f.write(file.file.read())
    att = models.Attachment(request_id=request_id, file_name=file.filename, file_path=path,
                             kind=kind, uploaded_by=user.id)
    db.add(att)
    log_action(db, user, "sample_request", request_id, "attach_file", file.filename)
    db.commit()
    return {"id": att.id, "file_name": att.file_name}
